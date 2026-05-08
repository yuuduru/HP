#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
布村伸一さん提供の Excel（季刊「教会」総目次）を data/issues.json に取り込む。

入力: ../Downloads/季刊「教会」総目次（20230626現在）.xlsx
出力: data/issues.json（バックアップは data/issues.backup.json）

各シート＝セクション種別、各行＝1記事。列は タイトル / 氏名 / 掲載号 / 備考。
xlsxにデータがある号は sections を xlsx ベースで再構築する。
"""

import json
import os
import shutil
from collections import OrderedDict
from pathlib import Path

import openpyxl

XLSX_PATH = Path(os.path.expanduser(
    "~/Downloads/季刊「教会」総目次（20230626現在）.xlsx"
))
DATA_FILE = Path(__file__).parent / "data" / "issues.json"
BACKUP_FILE = Path(__file__).parent / "data" / "issues.backup.json"

# xlsx シート名 → 出力セクション名（既存issues.jsonとの名寄せ）
SHEET_TO_SECTION = {
    "巻頭論壇": "巻頭言",
    "ＱＫ特集": "特集",
    "ＱK論文(神学研究)": "論文",
    "QK辛口誌上論争": "辛口誌上論争",
    "閑話休題 ティー・ブレイク": "ティー・ブレイク",
    "ＱＫ小論": "小論",
    "御言葉への沈潜": "御言葉への沈潜",
    "御言葉の味わい": "御言葉の味わい",
    "QK随想": "随想",
    "特": "特集",
    "ハイデルベルク信仰問答に基づくリレー黙想": "ハイデルベルク信仰問答リレー黙想",
    "ジュネーヴ教会信仰問答リレー黙想": "ジュネーヴ教会信仰問答リレー黙想",
    "ウェストミンスター小教理問答リレー黙想": "ウェストミンスター小教理問答リレー黙想",
    "ウェストミンスター大教理問答リレー黙想": "ウェストミンスター大教理問答リレー黙想",
    "本のオアシス": "本のオアシス",
    "聖書霊解(Lectio Spiritualis)": "聖霊解",
    "海外ニュース": "海外ニュース",
    "宗教改革五〇〇周年記念牧師会講師紹介": "宗教改革500周年記念牧師会講師紹介",
    "緊急報告 熊本地震": "緊急報告 熊本地震",
    "日曜学校の説教のために": "日曜学校の説教のために",
    "牧会話": "牧会話",
}

# セクション表示順（小さい数字ほど上に出す）
SECTION_ORDER = [
    "巻頭言",
    "特集",
    "論文",
    "辛口誌上論争",
    "小論",
    "随想",
    "聖霊解",
    "聖書研究",
    "聖書講義",
    "御言葉への沈潜",
    "御言葉の味わい",
    "翻訳",
    "牧会話",
    "牧会者の証言",
    "牧会者通信",
    "講座",
    "ティー・ブレイク",
    "緊急報告 熊本地震",
    "宗教改革500周年記念牧師会講師紹介",
    "日曜学校の説教のために",
    "ハイデルベルク信仰問答リレー黙想",
    "ジュネーヴ教会信仰問答リレー黙想",
    "ウェストミンスター小教理問答リレー黙想",
    "ウェストミンスター大教理問答リレー黙想",
    "本のオアシス",
    "海外ニュース",
]


def section_sort_key(name: str) -> tuple:
    if name in SECTION_ORDER:
        return (0, SECTION_ORDER.index(name))
    return (1, name)  # 未知セクションは末尾にアルファベット順


def parse_xlsx() -> dict:
    """xlsx を {issue_no: {section_name: [{title, author, note}, ...]}} に変換"""
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    issues: dict[int, dict[str, list]] = {}

    for sheet_name in wb.sheetnames:
        section_name = SHEET_TO_SECTION.get(sheet_name, sheet_name)
        ws = wb[sheet_name]

        # データ行は3行目以降（1行目=シート名, 2行目=ヘッダー）
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row:
                continue
            title = (row[0] or "").strip() if isinstance(row[0], str) else (row[0] if row[0] is not None else "")
            author = (row[1] or "").strip() if isinstance(row[1], str) else (row[1] if row[1] is not None else "")
            num = row[2] if len(row) > 2 else None
            note = (row[3] or "").strip() if len(row) > 3 and isinstance(row[3], str) else ""

            # 号数を整数化
            if isinstance(num, int):
                issue_no = num
            elif isinstance(num, float) and num.is_integer():
                issue_no = int(num)
            elif isinstance(num, str) and num.strip().isdigit():
                issue_no = int(num.strip())
            else:
                continue

            if not title and not author:
                continue

            article = {"title": str(title) if title else "[不明]",
                       "author": str(author) if author else "[不明]"}
            if note:
                article["note"] = note

            issues.setdefault(issue_no, OrderedDict()).setdefault(section_name, []).append(article)

    return issues


def merge_into_existing(xlsx_data: dict, existing: dict) -> dict:
    """既存 issues.json を xlsx データで上書き（xlsxにある号のみ）"""
    new_issues = []
    for issue in existing["issues"]:
        n = issue["number"]
        if n in xlsx_data:
            # xlsx の sections を SECTION_ORDER に従って並べる
            sections = []
            for sec_name in sorted(xlsx_data[n].keys(), key=section_sort_key):
                sections.append({
                    "sectionName": sec_name,
                    "articles": xlsx_data[n][sec_name],
                })
            new_issue = {
                "number": n,
                "date": issue.get("date", ""),
                "sections": sections,
            }
            if "label" in issue:
                new_issue["label"] = issue["label"]
            new_issues.append(new_issue)
        else:
            new_issues.append(issue)
    return {"issues": new_issues}


def main():
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"xlsx not found: {XLSX_PATH}")

    print(f"Loading: {XLSX_PATH}")
    xlsx_data = parse_xlsx()
    print(f"xlsx covers {len(xlsx_data)} issues: "
          f"{min(xlsx_data)}〜{max(xlsx_data)}")

    with open(DATA_FILE, encoding="utf-8") as f:
        existing = json.load(f)

    # バックアップ
    if not BACKUP_FILE.exists():
        shutil.copy(DATA_FILE, BACKUP_FILE)
        print(f"Backup: {BACKUP_FILE}")

    merged = merge_into_existing(xlsx_data, existing)

    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(merged, f, ensure_ascii=False, indent=2)
    print(f"Wrote: {DATA_FILE}")

    # 統計
    filled_count = sum(
        1 for i in merged["issues"]
        for s in i.get("sections", [])
        for a in s.get("articles", [])
        if a.get("title") not in ("[不明]", "[未入力]", "")
    )
    print(f"Total filled articles: {filled_count}")


if __name__ == "__main__":
    main()
